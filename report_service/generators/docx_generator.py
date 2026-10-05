"""
QuardCube Labs - Professional Business Microsoft Word (.docx) Report Generator
Generates publication-quality, structured, editable business reports.
"""

import io
import os
from datetime import datetime
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set background color of a docx table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding of a docx table cell in twentieths of a point (dxa)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tc_pr.append(tc_mar)


def render_chart_image_bytes(chart_data, chart_type='bar', title=""):
    """Render a Matplotlib chart and return image bytes for Word document embedding."""
    if not chart_data:
        return None

    fig, ax = plt.subplots(figsize=(6.2, 3.0), dpi=180)
    fig.patch.set_facecolor('#FFFFFF')
    ax.set_facecolor('#F8FAFC')

    primary_color = '#0D9488'
    accent_navy = '#0F172A'

    labels = [str(item.get('label') or item.get('name') or item.get('date') or '') for item in chart_data][:12]
    values = [float(item.get('value') or item.get('revenue') or item.get('total') or item.get('count') or 0) for item in chart_data][:12]

    if not labels or not values or sum(values) == 0:
        plt.close(fig)
        return None

    if chart_type in ['line', 'area']:
        ax.plot(labels, values, color=primary_color, marker='o', linewidth=2.5, markersize=5)
        if chart_type == 'area':
            ax.fill_between(range(len(labels)), values, color=primary_color, alpha=0.15)
        ax.grid(True, linestyle='--', alpha=0.5, color='#CBD5E1')
        plt.xticks(rotation=22, ha='right', fontsize=8)
    elif chart_type in ['pie', 'doughnut']:
        colors_list = ['#0D9488', '#0F172A', '#3B82F6', '#8B5CF6', '#EC4899', '#F59E0B', '#10B981', '#64748B']
        wedge_props = dict(width=0.45 if chart_type == 'doughnut' else 1.0, edgecolor='white', linewidth=1.5)
        ax.pie(values, labels=labels, autopct='%1.1f%%', startangle=135, colors=colors_list[:len(values)], textprops={'fontsize': 8, 'fontweight': 'bold'})
    else:
        bars = ax.bar(labels, values, color=primary_color, edgecolor='#0F766E', width=0.52)
        ax.grid(axis='y', linestyle='--', alpha=0.5, color='#CBD5E1')
        plt.xticks(rotation=20, ha='right', fontsize=8)

    if title:
        ax.set_title(title, fontsize=10.5, fontweight='bold', color=accent_navy, pad=10)

    plt.tight_layout()
    img_buf = io.BytesIO()
    plt.savefig(img_buf, format='png', bbox_inches='tight')
    plt.close(fig)
    img_buf.seek(0)
    return img_buf


def generate_docx_bytes(report_data: dict) -> bytes:
    """Generate a clean, professional, styled Microsoft Word (.docx) document."""
    doc = Document()

    # Set page margins to 0.75 in
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # 1. Company Header
    branding = report_data.get('branding', {})
    company_name = branding.get('companyName', 'QUARDCUBE LABS')
    company_sub = branding.get('subtitle', 'Enterprise Technology & Infrastructure Solutions')
    co_addr = branding.get('address', 'Makumbusho, Millennium Tower 14th Floor, Dar es Salaam, Tanzania')
    co_contact = f"{branding.get('phone', '+255 623 893 383')} | {branding.get('email', 'info@quardcubelabs.co.tz')}"

    hdr_p = doc.add_paragraph()
    hdr_p.paragraph_format.space_after = Pt(2)
    c_run = hdr_p.add_run(company_name.upper())
    c_run.bold = True
    c_run.font.size = Pt(13)
    c_run.font.color.rgb = RGBColor(13, 148, 136) # Teal

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(8)
    s_run = sub_p.add_run(f"{company_sub}\n{co_addr} • {co_contact}")
    s_run.font.size = Pt(8.5)
    s_run.font.color.rgb = RGBColor(100, 116, 139)

    # Divider line
    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_after = Pt(12)
    d_run = div_p.add_run("―" * 60)
    d_run.font.color.rgb = RGBColor(15, 23, 42)

    # 2. Report Title & Metadata
    report_title = report_data.get('title', 'Executive Management Report')
    report_desc = report_data.get('subtitle') or report_data.get('description') or 'Comprehensive Performance Assessment'
    period_from = report_data.get('period', {}).get('from', '')
    period_to = report_data.get('period', {}).get('to', '')
    period_str = f"Reporting Window: {period_from} to {period_to}" if period_from and period_to else "Reporting Window: All Historical Activity"
    prepared_by = branding.get('preparedBy') or report_data.get('auditSeal', {}).get('officer') or 'Senior Reporting Officer'

    t_p = doc.add_paragraph()
    t_p.paragraph_format.space_after = Pt(2)
    t_run = t_p.add_run(report_title)
    t_run.bold = True
    t_run.font.size = Pt(18)
    t_run.font.color.rgb = RGBColor(15, 23, 42)

    m_p = doc.add_paragraph()
    m_p.paragraph_format.space_after = Pt(14)
    m_run = m_p.add_run(f"{report_desc}\n{period_str} • Prepared by: {prepared_by} • Issued: {datetime.now().strftime('%d %B %Y')}")
    m_run.font.size = Pt(9.5)
    m_run.font.italic = True
    m_run.font.color.rgb = RGBColor(71, 85, 105)

    # 3. Executive Summary Callout Box
    narrative = report_data.get('narrative', {})
    exec_summary = narrative.get('executiveSummary') or narrative.get('overview') or report_data.get('summary', {}).get('executiveSummary')

    if exec_summary:
        sum_table = doc.add_table(rows=1, cols=1)
        sum_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sum_cell = sum_table.rows[0].cells[0]
        sum_cell.width = Inches(7.0)
        set_cell_background(sum_cell, "F8FAFC")
        set_cell_margins(sum_cell, top=140, bottom=140, left=200, right=200)

        sum_p = sum_cell.paragraphs[0]
        sum_p.paragraph_format.space_after = Pt(4)
        h_run = sum_p.add_run("EXECUTIVE SUMMARY & ANALYSIS\n")
        h_run.bold = True
        h_run.font.size = Pt(10)
        h_run.font.color.rgb = RGBColor(13, 148, 136)

        b_run = sum_p.add_run(exec_summary)
        b_run.font.size = Pt(9.5)
        b_run.font.color.rgb = RGBColor(30, 41, 59)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 4. Key Metrics KPI Table
    metrics = report_data.get('summary', {}).get('metrics', [])
    if metrics:
        h1 = doc.add_heading('Key Performance Indicators', level=1)
        h1.runs[0].font.color.rgb = RGBColor(15, 23, 42)
        h1.runs[0].font.size = Pt(13)

        cols = min(3, len(metrics))
        rows_needed = (len(metrics) + cols - 1) // cols
        kpi_table = doc.add_table(rows=rows_needed, cols=cols)
        kpi_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        for idx, m in enumerate(metrics):
            r_idx = idx // cols
            c_idx = idx % cols
            cell = kpi_table.cell(r_idx, c_idx)
            cell.width = Inches(7.0 / cols)
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            v_run = p.add_run(f"{m.get('value', '0')}\n")
            v_run.bold = True
            v_run.font.size = Pt(13)
            v_run.font.color.rgb = RGBColor(13, 148, 136)

            l_run = p.add_run(f"{m.get('label', '')}\n")
            l_run.bold = True
            l_run.font.size = Pt(8.5)
            l_run.font.color.rgb = RGBColor(15, 23, 42)

            d_run = p.add_run(m.get('description', ''))
            d_run.font.size = Pt(7.5)
            d_run.font.color.rgb = RGBColor(100, 116, 139)

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 5. Dynamic Sections (Charts and Tables with Intros)
    sections = report_data.get('sections', [])
    charts_dict = report_data.get('charts', {})
    tables_dict = report_data.get('tables', {})
    section_narratives = narrative.get('sectionNarratives', {})

    for sec in sections:
        if not sec.get('enabled', True):
            continue

        sec_type = sec.get('type')
        data_key = sec.get('dataKey', '')
        sec_title = sec.get('title', '')
        intro_text = sec.get('introNarrative') or section_narratives.get(data_key)

        # A. CHART SECTION
        if sec_type == 'chart' and data_key in charts_dict:
            c_data = charts_dict[data_key]
            h2 = doc.add_heading(sec_title or c_data.get('title', 'Analytical Chart'), level=2)
            h2.runs[0].font.color.rgb = RGBColor(15, 23, 42)

            chart_intro = c_data.get('introText') or intro_text
            if chart_intro:
                p_intro = doc.add_paragraph(chart_intro)
                p_intro.paragraph_format.space_after = Pt(6)
                p_intro.runs[0].font.size = Pt(9)
                p_intro.runs[0].font.color.rgb = RGBColor(71, 85, 105)

            labels = c_data.get('labels', [])
            values = c_data.get('values', [])
            chart_type_val = sec.get('chartType') or c_data.get('chartType', 'bar')

            items_payload = [{'label': l, 'value': v} for l, v in zip(labels, values)]
            chart_img_buf = render_chart_image_bytes(items_payload, chart_type=chart_type_val, title=c_data.get('title', ''))

            if chart_img_buf:
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.add_run().add_picture(chart_img_buf, width=Inches(6.0))
                doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # B. TABLE SECTION
        elif sec_type == 'table' and data_key in tables_dict:
            t_data = tables_dict[data_key]
            h2 = doc.add_heading(sec_title or t_data.get('title', 'Data Register'), level=2)
            h2.runs[0].font.color.rgb = RGBColor(15, 23, 42)

            table_intro = t_data.get('introText') or intro_text
            if table_intro:
                p_intro = doc.add_paragraph(table_intro)
                p_intro.paragraph_format.space_after = Pt(6)
                p_intro.runs[0].font.size = Pt(9)
                p_intro.runs[0].font.color.rgb = RGBColor(71, 85, 105)

            headers = t_data.get('headers', [])
            raw_rows = t_data.get('rows', [])

            if headers and raw_rows:
                table = doc.add_table(rows=len(raw_rows[:40]) + 1, cols=len(headers))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER

                # Format Header Row
                hdr_cells = table.rows[0].cells
                for idx, text in enumerate(headers):
                    hdr_cells[idx].text = str(text).upper()
                    set_cell_background(hdr_cells[idx], "0F172A") # Navy
                    set_cell_margins(hdr_cells[idx], top=80, bottom=80, left=100, right=100)
                    p = hdr_cells[idx].paragraphs[0]
                    p.runs[0].font.bold = True
                    p.runs[0].font.size = Pt(8)
                    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)

                # Data Rows
                for r_idx, row_data in enumerate(raw_rows[:40], start=1):
                    row_cells = table.rows[r_idx].cells
                    bg_color = "F1F5F9" if r_idx % 2 == 0 else "FFFFFF"
                    
                    for c_idx, val in enumerate(row_data):
                        val_str = str(val if val is not None else '')
                        row_cells[c_idx].text = val_str
                        set_cell_background(row_cells[c_idx], bg_color)
                        set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=100, right=100)
                        
                        p = row_cells[c_idx].paragraphs[0]
                        p.runs[0].font.size = Pt(8)
                        p.runs[0].font.color.rgb = RGBColor(30, 41, 59)
                        
                        is_numeric = any(char.isdigit() for char in val_str) and ('TZS' in val_str or val_str.replace(',', '').replace('.', '').isdigit())
                        if is_numeric or c_idx >= len(headers) - 1:
                            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

                doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 6. Observations & Recommendations
    observations = narrative.get('observations', [])
    recommendations = narrative.get('recommendations', [])

    if observations or recommendations:
        h1 = doc.add_heading('Strategic Observations & Next Steps', level=1)
        h1.runs[0].font.color.rgb = RGBColor(15, 23, 42)

        if observations:
            p_obs_h = doc.add_paragraph()
            r_obs_h = p_obs_h.add_run("Key Findings:")
            r_obs_h.bold = True
            r_obs_h.font.size = Pt(9.5)
            r_obs_h.font.color.rgb = RGBColor(15, 23, 42)

            for obs in observations:
                p_bullet = doc.add_paragraph(obs, style='List Bullet')
                p_bullet.paragraph_format.space_after = Pt(3)
                p_bullet.runs[0].font.size = Pt(9)
                p_bullet.runs[0].font.color.rgb = RGBColor(51, 65, 85)

        if recommendations:
            p_rec_h = doc.add_paragraph()
            p_rec_h.paragraph_format.space_before = Pt(6)
            r_rec_h = p_rec_h.add_run("Actionable Recommendations:")
            r_rec_h.bold = True
            r_rec_h.font.size = Pt(9.5)
            r_rec_h.font.color.rgb = RGBColor(13, 148, 136)

            for rec in recommendations:
                p_bullet = doc.add_paragraph(rec, style='List Bullet')
                p_bullet.paragraph_format.space_after = Pt(3)
                p_bullet.runs[0].font.size = Pt(9)
                p_bullet.runs[0].font.color.rgb = RGBColor(51, 65, 85)

    # Save to buffer
    doc_buffer = io.BytesIO()
    doc.save(doc_buffer)
    doc_buffer.seek(0)
    return doc_buffer.getvalue()
